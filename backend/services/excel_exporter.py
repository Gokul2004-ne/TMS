import os
import json
from datetime import datetime, timezone, timedelta, date
from typing import Optional, Dict, Any, List
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_
from models import ClaimTimeline, Event, Associate
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Standard Voice Time & Motion taxonomy steps matching 'Prev Entd Voice T&M.xlsx'
STANDARD_TM_STEPS = [
    ("Login to Citrix application.", "BVA", 0.05),
    ("Login to Athena software.", "BVA", 0.04),
    ("Search and open the assigned Claim ID in Athena.", "VA", 0.08),
    ("Review claim details, denial reason, notes history, and billing information.", "VA", 0.12),
    ("Perform pre-service/claim analysis to identify the next action.", "VA", 0.10),
    ("Open the Phone Payer tab in Athena.", "VA", 0.05),
    ("Verify payer information and available contact details.", "VA", 0.04),
    ("Login to portal.", "BVA", 0.07),
    ("Review any available EOB, Previous, or documents.", "VA", 0.08),
    ("A call is placed to the Insurance", "VA", 0.14),
    ("Discuss about the claim/appeal/Reprocess status with the insurance representative", "VA", 0.12),
    ("Take the information from the representative", "VA", 0.06),
    ("Update detailed call notes in Athena.", "VA", 0.08),
    ("Validate that all supporting information has been documented accurately.", "VA", 0.04),
    ("Save all changes and move to the next claim.", "BVA", 0.03),
]

PAYERS = ["UHC", "Medicare", "BCBS", "Aetna", "GWH-Cigna", "Humana"]

HEADERS = [
    "Date",
    "User Name ",
    "WQ Name ",
    "Payor Name ",
    "Claim #",
    "Steps",
    "Start time",
    "End time",
    "Total Duration",
    "Voice/Non Voice",
    "Value add/NVA",
    "Resolved by",
    "Account type",
    "User Experience",
    "Unique Id"
]


def _format_time_ist(dt: Optional[datetime]) -> str:
    if not dt:
        dt = datetime.now(timezone.utc)
    # Convert to IST (+5:30)
    ist_dt = dt + timedelta(hours=5, minutes=30)
    return ist_dt.strftime("%H:%M:%S")


def _format_duration(total_sec: int) -> str:
    total_sec = max(0, total_sec)
    h = total_sec // 3600
    m = (total_sec % 3600) // 60
    s = total_sec % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


async def export_claim_to_excel(
    db: AsyncSession,
    associate_id: str,
    claim_id: str,
    excel_path: Optional[str] = None
) -> str:
    """
    Appends or updates the completed claim in an Excel workbook matching
    the exact 'Prev Entd Voice T&M.xlsx' schema.
    Also produces a CSV copy for maximum compatibility.
    """
    if not excel_path:
        project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
        excel_path = os.path.join(project_root, "TMS_Completed_Claims.xlsx")

    # Fetch associate name
    assoc_stmt = select(Associate).where(Associate.id == associate_id)
    assoc_res = await db.execute(assoc_stmt)
    assoc = assoc_res.scalars().first()
    assoc_name = assoc.name if assoc else f"Associate {associate_id}"

    # Fetch claim timeline
    tl_stmt = select(ClaimTimeline).where(
        and_(ClaimTimeline.associate_id == associate_id, ClaimTimeline.claim_id == claim_id)
    )
    tl_res = await db.execute(tl_stmt)
    timeline = tl_res.scalars().first()

    total_sec = max(60, timeline.total_duration_seconds if timeline and timeline.total_duration_seconds else 300)
    base_start = timeline.start_time if timeline and timeline.start_time else datetime.now(timezone.utc)
    flags = json.loads(timeline.nva_flags_json if timeline and timeline.nva_flags_json else "[]")
    session_id = timeline.session_id if timeline and timeline.session_id else f"sess-{associate_id.lower()}-1"

    # Construct granular steps
    claim_steps = list(STANDARD_TM_STEPS)
    if "EXCEL_OVERUSE" in flags:
        claim_steps.insert(9, ("Verify manual calculation and adjustments in Excel.", "NVA", 0.10))
    if "APP_SWITCHING" in flags:
        claim_steps.insert(8, ("Cross-reference patient eligibility across clearinghouse portals.", "NVA", 0.08))
    if "LONG_IDLE" in flags:
        claim_steps.insert(7, ("Review policy guidelines; pause for supervisor override authorization.", "NVA", 0.12))

    weight_sum = sum(w for _, _, w in claim_steps)
    curr_time = base_start
    today_str = base_start.strftime("%Y-%m-%d") if base_start else date.today().isoformat()
    payer = PAYERS[hash(claim_id) % len(PAYERS)]

    rows_data: List[List[Any]] = []
    for name, cat, w in claim_steps:
        dur_sec = max(5, int((w / weight_sum) * total_sec))
        end_step = curr_time + timedelta(seconds=dur_sec)
        row = [
            today_str,
            assoc_name,
            "Prev Entd",
            payer,
            claim_id,
            name,
            _format_time_ist(curr_time),
            _format_time_ist(end_step),
            _format_duration(dur_sec),
            "Voice",
            cat,
            "Call",
            "Medium",
            "Tenured",
            session_id
        ]
        rows_data.append(row)
        curr_time = end_step

    # --- Write to openpyxl Workbook ---
    wb = None
    ws = None
    if os.path.exists(excel_path):
        try:
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active
        except Exception:
            wb = None

    if wb is None or ws is None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Completed Claims"

        # Executive Header styling (Obsidian Blue)
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.append(HEADERS)
        for col_num in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

    # Remove previous rows for this claim if already exists to avoid duplicates
    existing_rows = list(ws.iter_rows(values_only=False))
    if len(existing_rows) > 1:
        # Check rows from bottom up
        for r_idx in range(len(existing_rows), 1, -1):
            claim_cell = ws.cell(row=r_idx, column=5).value  # Column E is Claim #
            if claim_cell == claim_id:
                ws.delete_rows(r_idx, 1)

    # Append new rows
    data_font = Font(name="Calibri", size=10)
    border_cell = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row_values in rows_data:
        ws.append(row_values)
        row_num = ws.max_row
        for col_num in range(1, len(row_values) + 1):
            c = ws.cell(row=row_num, column=col_num)
            c.font = data_font
            c.border = border_cell
            if col_num in [1, 5, 7, 8, 9, 10, 11, 12, 13, 14]:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    try:
        wb.save(excel_path)
    except Exception as e:
        print(f"[Excel Export] Warning saving .xlsx: {e}")

    # Also update companion CSV
    csv_path = excel_path.replace(".xlsx", ".csv")
    try:
        import csv
        file_exists = os.path.exists(csv_path)
        with open(csv_path, mode="a" if file_exists else "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            if not file_exists:
                writer.writerow(HEADERS)
            for r in rows_data:
                writer.writerow(r)
    except Exception as e:
        print(f"[CSV Export] Warning saving .csv: {e}")

    return excel_path
